#!/usr/bin/env python3
"""Build a Chinese deck describing the current TP/AP experiment design.

The deck is intentionally based on the checked-in scenario definitions and
the current lab contract.  It explains the distinction between the complete
TPCC case catalog and the six scenarios exposed by the console, and between
the AP query catalog and a fixed pressure scenario.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parent
DEPS = ROOT / ".pptx_deps"
if DEPS.is_dir():
    sys.path.insert(0, str(DEPS))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402


OUTPUT = ROOT.parent / "TP_AP场景设计.pptx"
FONT = "Noto Sans CJK SC"

NAVY = RGBColor(15, 29, 55)
INK = RGBColor(28, 42, 66)
MUTED = RGBColor(86, 104, 130)
LIGHT = RGBColor(246, 248, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(216, 225, 237)
TP_BLUE = RGBColor(39, 112, 220)
TP_CYAN = RGBColor(22, 162, 181)
AP_PURPLE = RGBColor(119, 80, 190)
AP_ORANGE = RGBColor(227, 126, 39)
GREEN = RGBColor(29, 155, 108)
RED = RGBColor(210, 78, 78)
GOLD = RGBColor(211, 151, 38)
PALE_BLUE = RGBColor(232, 241, 255)
PALE_CYAN = RGBColor(230, 248, 248)
PALE_PURPLE = RGBColor(241, 235, 253)
PALE_ORANGE = RGBColor(255, 243, 226)
PALE_GREEN = RGBColor(230, 247, 239)
PALE_RED = RGBColor(253, 237, 237)
PALE_GOLD = RGBColor(255, 248, 224)


def set_run_font(run, size, color, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    try:
        from pptx.oxml.ns import qn

        run._r.get_or_add_rPr().set(qn("a:ea"), FONT)
    except Exception:
        pass


def add_text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=12,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.06,
    italic=False,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(value)
    set_run_font(run, size, color, bold, italic)
    return shape


def add_box(slide, x, y, w, h, fill, line=LINE, radius=True, width=0.7):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    return shape


def add_header(slide, number, title, subtitle):
    add_box(slide, 0, 0, 13.333, 1.02, NAVY, NAVY, radius=False)
    add_text(
        slide,
        f"TP × AP  SCENARIO DESIGN  /  {number}",
        0.52,
        0.11,
        4.0,
        0.18,
        size=7.3,
        color=RGBColor(159, 191, 236),
        bold=True,
        margin=0,
    )
    add_text(
        slide,
        title,
        0.52,
        0.30,
        12.0,
        0.40,
        size=23,
        color=WHITE,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )
    add_text(
        slide,
        subtitle,
        0.54,
        0.77,
        12.0,
        0.16,
        size=8.8,
        color=RGBColor(193, 207, 228),
        margin=0,
    )


def add_footer(slide, value):
    add_text(slide, value, 0.52, 7.22, 12.2, 0.15, size=6.6, color=MUTED, margin=0)


def add_pill(slide, text, x, y, w, color, pale, size=8.5):
    add_box(slide, x, y, w, 0.28, pale, pale)
    add_text(slide, text, x + 0.03, y + 0.025, w - 0.06, 0.22, size=size,
             color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_arrow(slide, x, y, w=0.22, h=0.32, color=RGBColor(177, 191, 211)):
    arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()
    return arrow


def add_flow_card(slide, x, y, w, h, number, title, body, accent, pale):
    add_box(slide, x, y, w, h, WHITE, LINE)
    add_box(slide, x, y, w, 0.08, accent, accent, radius=False)
    add_text(slide, number, x + 0.13, y + 0.15, 0.36, 0.20, size=8.3, color=accent, bold=True, margin=0)
    add_text(slide, title, x + 0.13, y + 0.39, w - 0.25, 0.27, size=13.2, color=INK, bold=True, margin=0)
    add_text(slide, body, x + 0.13, y + 0.78, w - 0.25, h - 0.88, size=8.8, color=MUTED, margin=0)


def add_metric_card(slide, x, y, w, value, label, accent, pale):
    add_box(slide, x, y, w, 1.00, WHITE, LINE)
    add_box(slide, x, y, 0.08, 1.00, accent, accent, radius=False)
    add_text(slide, value, x + 0.18, y + 0.14, w - 0.30, 0.39, size=22, color=accent, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, label, x + 0.16, y + 0.61, w - 0.28, 0.22, size=8.4, color=MUTED,
             align=PP_ALIGN.CENTER, margin=0)


def read_only_sql(text):
    text = re.sub(r"--[^\n]*", " ", text)
    text = re.sub(r"/\*.*?\*/", " ", text, flags=re.S)
    text = text.strip().rstrip(";").strip().lower()
    if not text or ";" in text or not text.startswith(("select", "with", "explain")):
        return False
    forbidden = r"\b(insert|update|delete|merge|create|drop|alter|truncate|grant|revoke|copy|call|do|vacuum|refresh|set|reset|begin|commit|rollback|prepare|execute|lock)\b"
    return re.search(forbidden, text) is None


def catalog_facts():
    """Read the current checked-in catalogs so the headline counts stay honest."""
    import importlib.util

    spec = importlib.util.spec_from_file_location("tpcc_transaction_cases", ROOT / "tpcc_transaction_cases.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    tp_cases = list(module.CASE_DEFINITIONS)
    focus = list(module.FOCUS_SCENARIO_IDS)
    sql_root = ROOT.parent / "dream" / "data" / "slow_queries" / "TPC-DS"
    sql_files = sorted(sql_root.glob("*.sql"), key=lambda p: int(p.stem) if p.stem.isdigit() else p.stem)
    read_only = sum(read_only_sql(p.read_text(encoding="utf-8")) for p in sql_files)
    return len(tp_cases), focus, len(sql_files), read_only


def build_deck():
    tp_count, focus_ids, ap_count, ap_read_only = catalog_facts()
    assert tp_count == 27
    assert len(focus_ids) == 6
    assert ap_count == 99
    assert ap_read_only == 95

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = "TP × AP 场景设计"
    prs.core_properties.subject = "SysInsight 与 DREAM 当前实验场景、链路和验证边界"
    prs.core_properties.author = "Codex"

    # 01. Cover and executive summary.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_box(slide, 0, 0, 13.333, 7.5, NAVY, NAVY, radius=False)
    add_box(slide, 0, 0, 13.333, 0.12, TP_BLUE, TP_BLUE, radius=False)
    add_text(slide, "SYSINSIGHT  ×  DREAM", 0.72, 0.75, 5.2, 0.25, size=10, color=RGBColor(160, 198, 248), bold=True, margin=0)
    add_text(slide, "TP × AP 场景设计", 0.68, 1.23, 8.8, 0.72, size=35, color=WHITE, bold=True, margin=0)
    add_text(slide, "当前实验控制台的工作负载、压力模型、调优链路与验证边界", 0.72, 2.10, 8.6, 0.28, size=15, color=RGBColor(193, 207, 228), margin=0)

    add_box(slide, 0.72, 3.02, 11.90, 0.78, RGBColor(28, 48, 82), RGBColor(53, 78, 119))
    add_text(slide, "一套 PostgreSQL · 两条工作负载链路 · 两种优化目标", 1.02, 3.22, 11.30, 0.30,
             size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)

    add_metric_card(slide, 0.72, 4.30, 3.55, "27", "TPCC 代码场景定义", TP_BLUE, PALE_BLUE)
    add_metric_card(slide, 4.88, 4.30, 3.55, "6", "控制台当前关注 TP 场景", TP_CYAN, PALE_CYAN)
    add_metric_card(slide, 9.04, 4.30, 3.58, "99 / 95", "TPC-DS SQL / 只读 SQL", AP_PURPLE, PALE_PURPLE)
    add_text(slide, "环境：PostgreSQL 12 · keeninsight · TPCC schema + tpcds schema · pgbench · pg_stat_statements · pg_hint_plan",
             0.72, 6.72, 12.0, 0.22, size=8.3, color=RGBColor(164, 183, 211), margin=0)
    add_text(slide, "2026-09-17  ·  场景设计说明", 0.72, 7.10, 4.0, 0.16, size=7.0, color=RGBColor(133, 158, 193), margin=0)

    # 02. Common architecture.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "02", "一条链路，两种目标", "TP 关注并发事务 TPS；AP 关注复杂分析 SQL 的实际耗时与改写质量")

    add_box(slide, 0.52, 1.24, 5.94, 4.94, WHITE, LINE)
    add_box(slide, 0.52, 1.24, 5.94, 0.09, TP_BLUE, TP_BLUE, radius=False)
    add_text(slide, "TP｜SysInsight 资源/参数恢复", 0.78, 1.49, 5.40, 0.28, size=16, color=INK, bold=True, margin=0)
    tp_flow = [
        ("01", "TPCC 正常负载", "tp_normal.sql\n五类 TPC-C 事务混合", TP_BLUE, PALE_BLUE),
        ("02", "外部 TP 压力", "固定速率 pgbench -R\n保持压力不撤掉", TP_CYAN, PALE_CYAN),
        ("03", "观测与检测", "Prometheus 告警 + perf\nSysInsight 原始检测", AP_ORANGE, PALE_ORANGE),
        ("04", "临时应用复测", "GPT 候选 → 临时配置\n压力下比较业务 TPS", GREEN, PALE_GREEN),
    ]
    x_positions = [0.78, 3.64]
    y_positions = [2.02, 3.58]
    for idx, (num, title, body, accent, pale) in enumerate(tp_flow):
        x = x_positions[idx % 2]
        y = y_positions[idx // 2]
        add_box(slide, x, y, 2.42, 1.15, pale, pale)
        add_text(slide, num, x + 0.13, y + 0.13, 0.32, 0.18, size=8.0, color=accent, bold=True, margin=0)
        add_text(slide, title, x + 0.13, y + 0.38, 2.12, 0.23, size=11.5, color=INK, bold=True, margin=0)
        add_text(slide, body, x + 0.13, y + 0.69, 2.13, 0.34, size=8.3, color=MUTED, margin=0)
    add_text(slide, "目标：压力仍在时，目标 TPCC 业务 TPS 回升；实验结束后配置恢复。", 0.78, 5.58, 5.35, 0.25,
             size=9.7, color=TP_BLUE, bold=True, margin=0)

    add_box(slide, 6.86, 1.24, 5.94, 4.94, WHITE, LINE)
    add_box(slide, 6.86, 1.24, 5.94, 0.09, AP_PURPLE, AP_PURPLE, radius=False)
    add_text(slide, "AP｜DREAM 单 SQL 优化", 7.12, 1.49, 5.40, 0.28, size=16, color=INK, bold=True, margin=0)
    ap_flow = [
        ("01", "选择 TPC-DS SQL", "Q1–Q99 目录\n单条只读 SELECT/WITH", AP_PURPLE, PALE_PURPLE),
        ("02", "原 SQL 基线", "真实执行并记录耗时\npg_stat_statements 观测", AP_ORANGE, PALE_ORANGE),
        ("03", "DREAM 异步分析", "诊断 → 计划/记忆检索\nGPT API 生成候选", TP_BLUE, PALE_BLUE),
        ("04", "优化 SQL 再执行", "rewrite / session / hint\n比较真实执行耗时", GREEN, PALE_GREEN),
    ]
    for idx, (num, title, body, accent, pale) in enumerate(ap_flow):
        x = [7.12, 9.98][idx % 2]
        y = [2.02, 3.58][idx // 2]
        add_box(slide, x, y, 2.42, 1.15, pale, pale)
        add_text(slide, num, x + 0.13, y + 0.13, 0.32, 0.18, size=8.0, color=accent, bold=True, margin=0)
        add_text(slide, title, x + 0.13, y + 0.38, 2.12, 0.23, size=11.5, color=INK, bold=True, margin=0)
        add_text(slide, body, x + 0.13, y + 0.69, 2.13, 0.34, size=8.3, color=MUTED, margin=0)
    add_text(slide, "目标：同一条 SQL 的 optimized 实测耗时低于 original；候选是否发布由验证策略决定。", 7.12, 5.58, 5.35, 0.25,
             size=9.7, color=AP_PURPLE, bold=True, margin=0)
    add_box(slide, 0.52, 6.42, 12.28, 0.51, NAVY, NAVY)
    add_text(slide, "共用底座：bridge 负责自动编排与审计，实验控制台负责触发，Grafana 负责趋势，PostgreSQL 负责真实执行。",
             0.76, 6.54, 11.82, 0.24, size=10.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_footer(slide, "设计原则：TP 的压力不能随数据库变慢而自动降速；AP 的预测收益不能替代数据库实际 replay。")

    # 03. TP catalog map.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "03", "TP 场景库全貌｜27 个定义，6 个控制台焦点", "代码保留完整探索空间；控制台只展示当前用于稳定演示和复测的焦点集合")

    groups = [
        ("基础资源压力", "3", "WAL / checkpoint\nPayment 提交\nautovacuum churn", TP_BLUE, PALE_BLUE),
        ("单事务突发", "4", "New Order\nDelivery\nOrder Status / Stock Level", TP_CYAN, PALE_CYAN),
        ("跨仓访问", "2", "Remote Payment\nRemote New Order", AP_PURPLE, PALE_PURPLE),
        ("热点访问", "4", "热点 New Order\n热点 Payment\n热点读事务", AP_ORANGE, PALE_ORANGE),
        ("混合事务", "2", "读混合\n五类 TPCC 混合", GREEN, PALE_GREEN),
        ("强度/复合变体", "12", "中等/高强度\n随机 I/O\n远程/热点/混合高压", GOLD, PALE_GOLD),
    ]
    positions = [(0.52, 1.36), (4.43, 1.36), (8.34, 1.36), (0.52, 3.36), (4.43, 3.36), (8.34, 3.36)]
    for (title, count, body, accent, pale), (x, y) in zip(groups, positions):
        add_box(slide, x, y, 3.45, 1.58, WHITE, LINE)
        add_box(slide, x, y, 0.11, 1.58, accent, accent, radius=False)
        add_text(slide, count, x + 0.28, y + 0.20, 0.58, 0.45, size=25, color=accent, bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, title, x + 1.02, y + 0.23, 2.13, 0.24, size=12.5, color=INK, bold=True, margin=0)
        add_text(slide, body, x + 1.02, y + 0.61, 2.16, 0.70, size=9.0, color=MUTED, margin=0)
    add_box(slide, 0.52, 5.31, 12.28, 1.16, NAVY, NAVY)
    add_text(slide, "控制台当前 6 个关注场景", 0.82, 5.54, 2.35, 0.24, size=12.5, color=WHITE, bold=True, margin=0)
    focus_text = "1 Order Status 突发　·　2 Stock Level 突发　·　3 Payment 高强度　·　4 Order Status 随机读 I/O　·　5 热点 Stock Level　·　6 热点 Payment"
    add_text(slide, focus_text, 3.10, 5.48, 9.24, 0.40, size=10.0, color=RGBColor(218, 229, 246), margin=0)
    add_text(slide, "注意：27 个是代码场景库；“当前可从实验控制台直接选择”指上面 6 个 focus 场景。其余定义仍可由 runner 复现，但不是控制台默认焦点。",
             0.78, 6.72, 11.84, 0.22, size=7.8, color=MUTED, margin=0)
    add_footer(slide, "来源：perf-anomaly-demo/tpcc_transaction_cases.py；控制台由 FOCUS_SCENARIO_IDS 过滤场景。")

    # 04. TP focus matrix.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "04", "TP 控制台焦点｜6 组外部压力如何制造退化", "外部压力使用固定速率 pgbench -R；业务 TPS 使用 tp_normal.sql 作为被保护目标")

    focus = [
        ("01", "Order Status 突发", "6600", "客服随机索引读并发", "blks_read / hit · active_total", TP_BLUE, PALE_BLUE),
        ("02", "Stock Level 突发", "3800", "最近订单行 + 库存范围检查", "blks_read / hit · buffers_backend", TP_CYAN, PALE_CYAN),
        ("03", "Payment 高强度", "3350", "提交与 WAL 写入竞争", "xact_commit · checkpoint · fsync", AP_ORANGE, PALE_ORANGE),
        ("04", "Order Status 随机读 I/O", "5600", "更强随机 I/O 访问压力", "blks_read / hit · active_total", AP_PURPLE, PALE_PURPLE),
        ("05", "热点 Stock Level", "3800", "同一仓库/地区读热点", "blks_read / hit · active_total", GREEN, PALE_GREEN),
        ("06", "热点 Payment", "1950", "同一仓库/地区更新热点", "lock_waits · xact_commit · fsync", RED, PALE_RED),
    ]
    headers = ["焦点", "场景", "外部目标 TPS", "压力画像", "主要观测指标"]
    widths = [0.75, 2.42, 1.60, 3.38, 3.72]
    x0 = 0.52
    cx = x0
    for header, width in zip(headers, widths):
        add_box(slide, cx, 1.30, width, 0.42, NAVY, NAVY, radius=False)
        add_text(slide, header, cx + 0.04, 1.39, width - 0.08, 0.20, size=8.4, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        cx += width + 0.02
    y = 1.77
    for idx, (num, title, target, image, evidence, accent, pale) in enumerate(focus):
        fill = WHITE if idx % 2 == 0 else RGBColor(242, 246, 251)
        cx = x0
        vals = [num, title, target, image, evidence]
        for col, (value, width) in enumerate(zip(vals, widths)):
            add_box(slide, cx, y, width, 0.64, fill, LINE, radius=False)
            if col == 0:
                add_text(slide, value, cx, y + 0.19, width, 0.22, size=11.5, color=accent, bold=True,
                         align=PP_ALIGN.CENTER, margin=0)
            elif col == 2:
                add_text(slide, value, cx, y + 0.19, width, 0.22, size=12, color=accent, bold=True,
                         align=PP_ALIGN.CENTER, margin=0)
            else:
                add_text(slide, value, cx + 0.10, y + 0.13, width - 0.18, 0.35, size=8.8,
                         color=INK if col == 1 else MUTED, bold=(col == 1), margin=0)
            cx += width + 0.02
        y += 0.69

    add_box(slide, 0.52, 6.10, 12.28, 0.72, PALE_BLUE, PALE_BLUE)
    add_text(slide, "指标口径", 0.78, 6.27, 0.95, 0.20, size=9.8, color=TP_BLUE, bold=True, margin=0)
    add_text(slide, "外部目标 TPS 只表示施压强度；真正要保护的是 tp_normal.sql 的目标业务 TPS。调优前后保持同等压力，才能把 TPS 回升归因于 SysInsight。",
             1.78, 6.20, 10.62, 0.34, size=9.0, color=INK, margin=0)
    add_footer(slide, "当前固定速率目标来自 lab_controller.py；旧闭环实验的外部压力 TPS 不直接与新实验混比。")

    # 05. TP lifecycle.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "05", "TP 实验时序｜先建立退化，再在压力中验证恢复", "当前控制台采用“基线 → 持续压力”两阶段，不再增加撤压后的自然恢复阶段")

    phases = [
        ("基线", "60 s", "tp_normal.sql\n记录稳态业务 TPS", TP_BLUE, PALE_BLUE),
        ("压力前半段", "持续压力", "启动外部 pgbench -R\n等待告警 / perf / 检测", RED, PALE_RED),
        ("调优后半段", "持续压力", "临时应用候选\n同压力复测业务 TPS", GREEN, PALE_GREEN),
        ("结束", "自动清理", "TemporaryPostgresConfiguration\n恢复配置、结束进程", AP_PURPLE, PALE_PURPLE),
    ]
    x = 0.52
    for i, (title, duration, body, accent, pale) in enumerate(phases):
        add_box(slide, x, 1.55, 2.75, 1.62, WHITE, LINE)
        add_box(slide, x, 1.55, 2.75, 0.10, accent, accent, radius=False)
        add_text(slide, title, x + 0.15, 1.83, 1.90, 0.25, size=14, color=INK, bold=True, margin=0)
        add_pill(slide, duration, x + 1.92, 1.80, 0.64, accent, pale, size=7.2)
        add_text(slide, body, x + 0.15, 2.28, 2.42, 0.55, size=9.0, color=MUTED, margin=0)
        if i < len(phases) - 1:
            add_arrow(slide, x + 2.87, 2.16)
        x += 3.08

    add_box(slide, 0.52, 3.65, 5.90, 2.42, WHITE, LINE)
    add_text(slide, "SysInsight 在 TP 链路里做什么", 0.80, 3.91, 5.28, 0.24, size=14, color=INK, bold=True, margin=0)
    sys_items = [
        "Prometheus 告警确认：系统确实进入异常窗口",
        "perf + 原始异常函数提取：定位热点函数/资源关联",
        "GPT API 解析候选：从多个候选配置中选择方案",
        "临时 session 配置应用：候选结束后恢复，不写死生产配置",
    ]
    add_text(slide, "\n".join("· " + item for item in sys_items), 0.82, 4.34, 5.23, 1.40, size=9.1, color=MUTED, margin=0)

    add_box(slide, 6.90, 3.65, 5.90, 2.42, WHITE, LINE)
    add_text(slide, "TP 成功判定", 7.18, 3.91, 5.28, 0.24, size=14, color=INK, bold=True, margin=0)
    pass_items = [
        "外部压力前后目标 TPS 在允许偏差内",
        "压力阶段目标业务 TPS 明显下降",
        "应用候选后，在同等压力下目标 TPS 回升",
        "应用状态为 applied_and_restored，配置恢复成功",
    ]
    add_text(slide, "\n".join("· " + item for item in pass_items), 7.20, 4.34, 5.23, 1.40, size=9.1, color=MUTED, margin=0)
    add_footer(slide, "实验结束时不会保留外部压力；但“是否调优有效”的证据只看压力仍存在时的目标业务 TPS。")

    # 06. AP design and data caveat.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "06", "AP 场景现状｜99 条 TPC-DS SQL，按单条查询进入 DREAM", "AP 当前是“查询驱动”场景，不是像 TP 那样预设 6 组外部压力配方")

    add_metric_card(slide, 0.52, 1.30, 2.55, "Q1–Q99", "TPC-DS SQL 目录", AP_PURPLE, PALE_PURPLE)
    add_metric_card(slide, 3.28, 1.30, 2.55, "95", "允许单条执行的只读 SQL", AP_ORANGE, PALE_ORANGE)
    add_metric_card(slide, 6.04, 1.30, 2.55, "Q4", "当前代码标记的推荐入口", GREEN, PALE_GREEN)
    add_metric_card(slide, 8.80, 1.30, 3.99, "Q8 ≠ 有效案例", "当前空数据状态下不能证明优化收益", RED, PALE_RED)

    add_box(slide, 0.52, 2.62, 7.05, 3.55, WHITE, LINE)
    add_text(slide, "AP 单 SQL 实验流程", 0.82, 2.90, 6.35, 0.24, size=14, color=INK, bold=True, margin=0)
    ap_steps = [
        ("1", "选择 SQL", "控制台从 99 条目录中选择一条只读 TPC-DS SQL", AP_PURPLE, PALE_PURPLE),
        ("2", "发送原 SQL", "真实执行一次，记录 original SQL 的耗时和执行信息", AP_ORANGE, PALE_ORANGE),
        ("3", "DREAM 异步分析", "诊断、计划分析、记忆检索和 GPT API 生成候选", TP_BLUE, PALE_BLUE),
        ("4", "再次执行优化结果", "执行 rewrite SQL、session 建议或显式允许的 Hint", GREEN, PALE_GREEN),
    ]
    y = 3.38
    for num, title, body, accent, pale in ap_steps:
        add_box(slide, 0.82, y, 0.48, 0.48, pale, pale)
        add_text(slide, num, 0.82, y + 0.12, 0.48, 0.20, size=10.5, color=accent, bold=True, align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, title, 1.52, y + 0.03, 1.48, 0.22, size=10.5, color=INK, bold=True, margin=0)
        add_text(slide, body, 3.10, y + 0.03, 4.12, 0.32, size=8.6, color=MUTED, margin=0)
        y += 0.62

    add_box(slide, 7.82, 2.62, 4.98, 3.55, PALE_RED, RGBColor(246, 202, 202))
    add_text(slide, "当前 AP 验证边界", 8.12, 2.90, 4.38, 0.24, size=14, color=RED, bold=True, margin=0)
    add_text(slide, "当前 tpcds schema 的事实表是空的。\n\n因此：\n· Q8 original 39.79 ms → optimized 39.46 ms，约 +0.83%\n· SQL 本体执行不到 1 ms，主要耗时来自连接/进程开销\n· DREAM 预测的 98% 不是数据库 replay 的实际收益\n\n要得到可信 AP 效果，必须先装载 TPC-DS 数据、ANALYZE，再用多次执行的中位数筛选案例。",
             8.12, 3.38, 4.35, 2.42, size=8.9, color=INK, margin=0)
    add_footer(slide, "结论：AP 链路已接通，但当前数据状态不适合把 Q8 或任意 Q1–Q99 查询称为“优化效果好”。")

    # 07. Comparison and UI.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "07", "TP 与 AP 对照｜同一平台，不同观测对象", "理解两类场景的差异，才能正确解释 Grafana 和实验控制台中的结果")

    x0 = 0.52
    widths = [2.15, 3.45, 3.45, 3.23]
    headers = ["维度", "TP｜SysInsight", "AP｜DREAM", "前端应该展示什么"]
    cx = x0
    for header, width in zip(headers, widths):
        add_box(slide, cx, 1.30, width, 0.43, NAVY, NAVY, radius=False)
        add_text(slide, header, cx + 0.05, 1.40, width - 0.10, 0.20, size=8.7, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        cx += width + 0.02
    rows = [
        ("工作负载", "TPCC 五类事务 + 外部 TPCC 压力", "TPC-DS 单条复杂只读 SQL", "场景/SQL 名称、压力目标"),
        ("主指标", "目标业务 TPS", "original / optimized 耗时 ms", "同一图上显示调优前后"),
        ("触发方式", "告警 → perf → SysInsight", "慢 SQL / 手动发送 → DREAM", "队列、阶段、耗时、错误"),
        ("优化对象", "PostgreSQL 参数/会话设置", "rewrite SQL / session / pg_hint_plan", "候选状态、应用状态"),
        ("安全边界", "临时应用、结束恢复", "只读校验、Hint 显式发布", "审计时间线与恢复结果"),
        ("当前限制", "部分历史场景需重新复测", "TPC-DS 事实表为空", "明确展示数据有效性警告"),
    ]
    y = 1.78
    for idx, row in enumerate(rows):
        fill = WHITE if idx % 2 == 0 else RGBColor(242, 246, 251)
        cx = x0
        for col, (value, width) in enumerate(zip(row, widths)):
            add_box(slide, cx, y, width, 0.70, fill, LINE, radius=False)
            add_text(slide, value, cx + 0.10, y + 0.12, width - 0.18, 0.42, size=8.65,
                     color=INK if col == 0 else MUTED, bold=(col == 0), margin=0)
            cx += width + 0.02
        y += 0.75

    add_box(slide, 0.52, 6.48, 12.28, 0.40, NAVY, NAVY)
    add_text(slide, "实验控制台：/lab　　Grafana：趋势与告警　　bridge：动作、队列、SQL 观测、应用/回滚审计",
             0.78, 6.57, 11.80, 0.20, size=9.4, color=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_footer(slide, "展示重点：把“系统发生了什么、哪个环节做了什么、实际是否有效”放在同一条时间线上。")

    # 08. Conclusions and next steps.
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(slide, "08", "当前设计结论与下一步", "先区分“链路已实现”和“性能案例有效”，再决定哪些场景进入正式演示")

    add_box(slide, 0.52, 1.33, 5.94, 4.70, WHITE, LINE)
    add_box(slide, 0.52, 1.33, 5.94, 0.10, TP_BLUE, TP_BLUE, radius=False)
    add_text(slide, "TP｜当前可用设计", 0.82, 1.64, 5.30, 0.28, size=16, color=INK, bold=True, margin=0)
    add_pill(slide, "链路完整", 4.92, 1.64, 1.15, GREEN, PALE_GREEN, size=7.7)
    tp_conclusion = [
        "代码场景库 27 个，控制台当前聚焦 6 个",
        "正常业务是 tp_normal.sql；外部压力为固定速率 pgbench",
        "SysInsight 负责告警、perf、候选配置和压力下复测",
        "TemporaryPostgresConfiguration 负责临时应用与恢复",
        "优先展示压力下 TPS 前后对比，不把外部 TPS 当业务 TPS",
    ]
    add_text(slide, "\n".join("· " + item for item in tp_conclusion), 0.84, 2.20, 5.12, 2.45, size=10.0, color=MUTED, margin=0)
    add_box(slide, 0.82, 4.96, 5.28, 0.66, PALE_BLUE, PALE_BLUE)
    add_text(slide, "演示建议：使用当前 6 个 focus 场景；历史达标结果仍需在当前状态复测后再做结论。",
             1.03, 5.12, 4.86, 0.27, size=8.9, color=TP_BLUE, bold=True, margin=0)

    add_box(slide, 6.86, 1.33, 5.94, 4.70, WHITE, LINE)
    add_box(slide, 6.86, 1.33, 5.94, 0.10, AP_PURPLE, AP_PURPLE, radius=False)
    add_text(slide, "AP｜当前链路已接通，数据需补齐", 7.16, 1.64, 5.30, 0.28, size=16, color=INK, bold=True, margin=0)
    add_pill(slide, "需先恢复数据", 11.28, 1.64, 1.20, RED, PALE_RED, size=7.0)
    ap_conclusion = [
        "SQL 目录 99 条，其中 95 条可单条只读执行",
        "DREAM 已支持原 SQL → 异步分析 → optimized replay",
        "Q4 是当前代码标记的推荐入口，Q8 当前不具备有效收益证明",
        "当前 tpcds 事实表为空，执行时间被连接开销主导",
        "下一步：装载 TPC-DS SF1 → ANALYZE → 重复测量 → 筛选案例",
    ]
    add_text(slide, "\n".join("· " + item for item in ap_conclusion), 7.18, 2.20, 5.12, 2.45, size=10.0, color=MUTED, margin=0)
    add_box(slide, 7.16, 4.96, 5.28, 0.66, PALE_PURPLE, PALE_PURPLE)
    add_text(slide, "筛选标准：结果正确 + 多次运行中位数稳定 + 实测收益，而不是只看 DREAM 预测值。",
             7.37, 5.12, 4.86, 0.27, size=8.9, color=AP_PURPLE, bold=True, margin=0)

    add_box(slide, 0.52, 6.35, 12.28, 0.59, NAVY, NAVY)
    add_text(slide, "最终目标：TP 证明系统在并发压力中可恢复；AP 证明复杂 SQL 在真实数据上可加速；两者都要由实际数据库执行结果闭环。",
             0.78, 6.50, 11.80, 0.25, size=10.4, color=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_footer(slide, "文件依据：tpcc_transaction_cases.py、lab_controller.py、TPCC场景验证报告.md、DREAM TPC-DS SQL catalog。")

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_deck()
    print(f"created {output}")
